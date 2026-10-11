"""Reading Excel workbooks (.xlsx): one worksheet, through openpyxl.

openpyxl is an optional dependency -- the ``excel`` extra -- so importing this
module never imports it; the import happens when an ``.xlsx`` file is actually
read, and its absence is a :class:`MissingDependencyError` naming the extra,
not a traceback.

An Excel workbook is not one table; it is a file of worksheets. So the rules
here are strict where the delimited and JSON readers are:

- The file must be a real ``.xlsx`` (a ZIP with an ``xl/workbook.xml`` member)
  that openpyxl can open. Anything else -- including an old-format ``.xls`` or
  a corrupt archive -- is :class:`MalformedInputError`, never a bare openpyxl
  or zip error.
- A workbook carrying VBA macros (an ``xl/vbaProject.bin`` member) is refused:
  a ``.xlsx`` is a data container, and macro-enabled content is out of the
  trust this library extends to an input file.
- Eligible relations are worksheets, including hidden and very-hidden ones.
  Chartsheets are not worksheets and never count as a second table.
- A workbook with more than one eligible worksheet is ambiguous: choosing a
  sheet would be a guess. Pass ``sheet=`` to name the one to read. The name is
  matched exactly. One eligible worksheet is selected automatically.
- Every successful read records that worksheet as ``source_selector`` with
  ``kind`` ``sheet`` and the exact title.
- The first row of the sheet is the header. Duplicate header cells are refused
  (``DuplicateHeadersError``), and an empty cell in the header is named by its
  column position so the file can be fixed.
- Formulas are read as their cached values (``data_only=True``); a formula
  whose value was never computed reads as missing, which the profiling layer
  reports like any other missing cell. The formula text is never executed.
- The row, column and cell budgets are enforced from the selected worksheet's
  declared dimensions **before** its cells are materialized. ``max_field_length``
  is enforced while those cells are scanned, before a DataFrame is built.
  An unselected worksheet is not measured against those logical limits.
"""

from __future__ import annotations

import importlib.util
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd

from aidatasetkit.core.exceptions import (
    CellLimitError,
    ColumnLimitError,
    DuplicateHeadersError,
    EmptyInputError,
    FieldLengthLimitError,
    InvalidIngestionOptionsError,
    MalformedInputError,
    MissingDependencyError,
    RowLimitError,
)
from aidatasetkit.ingestion.types import (
    IngestionLimits,
    LoadOptions,
    SourceKind,
    SourceSelector,
    TableFormat,
    _cells_exceed,
)

__all__ = ["read_excel"]


def read_excel(
    path: Path,
    fmt: TableFormat,
    options: LoadOptions,
    limits: IngestionLimits,
    sheet: "str | None" = None,
) -> tuple[pd.DataFrame, SourceSelector]:
    """The frame held by one worksheet of an .xlsx file, and which worksheet.

    Raises:
        InvalidIngestionOptionsError, MissingDependencyError, EmptyInputError,
        MalformedInputError, DuplicateHeadersError, RowLimitError,
        ColumnLimitError, CellLimitError, FieldLengthLimitError.
    """
    if not isinstance(sheet, str) and sheet is not None:
        raise InvalidIngestionOptionsError(
            f"sheet must name a worksheet, got {type(sheet).__name__}."
        )
    if not options.is_default:
        raise InvalidIngestionOptionsError(
            "LoadOptions (encoding, delimiter, header) apply to text files; "
            f"{path.name} is an Excel file, which describes itself."
        )
    _openpyxl(path)
    wb = _open(path)
    try:
        ws = _choose_sheet(path, wb, sheet)
        _check_dimensions(path, ws, limits)
        rows = _rows(ws, path)
        _reject_long_fields(path, rows, limits)
    finally:
        wb.close()
    return _frame(path, rows), _sheet_selector(sheet, ws)


def _openpyxl(path: Path):
    """Import openpyxl, or refuse naming the extra that provides it."""
    try:
        spec = importlib.util.find_spec("openpyxl")
    except ImportError:
        spec = None
    if spec is None:
        raise MissingDependencyError(
            "Reading .xlsx files needs openpyxl, which is not installed. It is "
            "an optional dependency: pip install 'aidatasetkit[excel]'."
        )


def _open(path: Path):
    """The workbook, after the archive itself is shown to be a macro-free xlsx."""
    import openpyxl

    if not zipfile.is_zipfile(path):
        raise MalformedInputError(
            f"{path.name} is not a readable .xlsx file (not a ZIP archive)."
        )
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
    if "xl/workbook.xml" not in names:
        raise MalformedInputError(
            f"{path.name} is not a readable .xlsx file (no workbook part)."
        )
    if "xl/vbaProject.bin" in names:
        raise MalformedInputError(
            f"{path.name} carries VBA macros. A .xlsx is a data container; "
            "macro-enabled content is refused. Save the data as a plain .xlsx."
        )
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise MalformedInputError(
            f"{path.name} is not a readable .xlsx file ({type(error).__name__})."
        ) from None


def _eligible_worksheets(wb) -> list:
    """Worksheets only. Hidden and veryHidden stay eligible. Chartsheets do not."""
    return [ws for ws in wb.worksheets if _worksheet_is_eligible(ws)]


def _worksheet_is_eligible(ws) -> bool:
    state = getattr(ws, "sheet_state", "visible")
    if state == "hidden":
        return True
    if state == "veryHidden":
        return True
    return True


def _choose_sheet(path: Path, wb, sheet: "str | None"):
    """The one worksheet to read: named, or the only one, else refused as ambiguous."""
    eligible = _eligible_worksheets(wb)
    names = [ws.title for ws in eligible]
    if sheet is not None:
        return _named_worksheet(path, wb, eligible, names, sheet)
    if len(names) == 0:
        raise EmptyInputError(
            f"{path.name} has no worksheet; there is no table to load."
        )
    if len(names) == 1:
        return eligible[0]
    raise MalformedInputError(
        f"{path.name} has {len(names)} sheets ({', '.join(names)}). Choosing one "
        "would be a guess; pass sheet= to name it."
    )


def _named_worksheet(path: Path, wb, eligible, names: list[str], sheet: str):
    """Exact title match. No case folding, trimming, or Unicode normalization."""
    for ws in eligible:
        if ws.title == sheet:
            return ws
    for chart in wb.chartsheets:
        if chart.title == sheet:
            raise MalformedInputError(
                f"{path.name} has no sheet named {sheet!r}. "
                f"{sheet!r} is a chartsheet, not a worksheet."
            )
    raise MalformedInputError(
        f"{path.name} has no sheet named {sheet!r}. Sheets: {', '.join(names)}."
    )


def _sheet_selector(requested: "str | None", chosen) -> SourceSelector:
    """The worksheet that was actually read, for evidence and config identity."""
    recorded = chosen.title
    if requested is None:
        recorded = chosen.title  # automatic selection is recorded too
    return SourceSelector(kind="sheet", name=recorded)


def _check_dimensions(path: Path, ws, limits: IngestionLimits) -> None:
    """Empty, row, column and cell budgets from the declared dimensions."""
    rows = ws.max_row or 0
    columns = ws.max_column or 0
    # A sheet with a header row but no data rows has one declared row; the data
    # rows are rows - 1. An empty sheet declares 1x1 with no cells at all.
    if rows == 0 or columns == 0 or (rows == 1 and columns == 1 and ws["A1"].value is None):
        raise EmptyInputError(
            f"{path.name} has no data; there is no table to load."
        )
    if limits.max_rows is not None and rows - 1 > limits.max_rows:
        raise RowLimitError(
            f"{path.name} exceeds the row limit ({rows - 1} data rows > {limits.max_rows}).",
            limit_name="max_rows", configured_limit=limits.max_rows,
            observed_value=rows - 1, input_kind=SourceKind.FILE.value,
        )
    if limits.max_columns is not None and columns > limits.max_columns:
        raise ColumnLimitError(
            f"{path.name} exceeds the column limit ({columns} > {limits.max_columns}).",
            limit_name="max_columns", configured_limit=limits.max_columns,
            observed_value=columns, input_kind=SourceKind.FILE.value,
        )
    if limits.max_cells is not None and _cells_exceed(rows, columns, limits.max_cells):
        cells = rows * columns
        raise CellLimitError(
            f"{path.name} exceeds the cell limit ({cells} > {limits.max_cells}).",
            limit_name="max_cells", configured_limit=limits.max_cells,
            observed_value=cells, input_kind=SourceKind.FILE.value,
        )


def _rows(ws, path: Path) -> list[tuple[Any, ...]]:
    """Every row's values, materialized only after the dimension budgets pass."""
    return [tuple(row) for row in ws.iter_rows(values_only=True)]


def _over_field_length(value: Any, limit: int) -> bool:
    """True when a textual or binary cell exceeds the limit.

    Numbers, booleans and datetimes are not measured by turning them into text.
    """
    if isinstance(value, str) or isinstance(value, (bytes, bytearray)):
        return len(value) > limit
    return False


def _reject_long_fields(path: Path, rows: list[tuple[Any, ...]], limits: IngestionLimits) -> None:
    """Refuse an oversized selected cell before a DataFrame is built."""
    if limits.max_field_length is None:
        return
    limit = limits.max_field_length
    for row in rows:
        for value in row:
            if _over_field_length(value, limit):
                length = len(value)
                raise FieldLengthLimitError(
                    f"{path.name} exceeds the field length limit ({length} > {limit}).",
                    limit_name="max_field_length",
                    configured_limit=limit,
                    observed_value=length,
                    input_kind=SourceKind.FILE.value,
                )


def _frame(path: Path, rows: list[tuple[Any, ...]]) -> pd.DataFrame:
    """Header + data rows -> a frame. First row is the header."""
    if not rows:
        raise EmptyInputError(f"{path.name} has no rows; there is no table to load.")
    header = list(rows[0])
    data = list(rows[1:])
    # A blank data row is a row of missing cells, not a row to drop: the
    # dimension preflight already counted it, and omitting it would be a silent
    # shorter table. A header with no data rows is the same empty input every
    # other reader refuses.
    if not data:
        raise EmptyInputError(
            f"{path.name} has a header row and no data rows; there is nothing to audit."
        )

    names: list[str] = []
    seen: set[str] = set()
    for index, cell in enumerate(header):
        name = str(cell) if cell is not None else f"__column_{index}"
        if cell is not None and name in seen:
            raise DuplicateHeadersError(
                f"{path.name} repeats the column label {name!r}. Give each column "
                "its own name first."
            )
        seen.add(name)
        names.append(name)
    return pd.DataFrame(data, columns=names)
