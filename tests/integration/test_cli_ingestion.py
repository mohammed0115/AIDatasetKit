"""G1-W1 through the command a user types: every file is read right or refused.

Each test runs the real CLI in a subprocess. A refused file must leave exactly
one line on stderr, exit 1, and publish nothing -- no ``CURRENT``, no run, and
an earlier good run left as it was.
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest

from aidatasetkit.cli import EXIT_CODES
from aidatasetkit.evidence import read_current

HEADER = ["tenure", "plan", "city", "churn"]
ROWS = [[str(3 + (i * 7) % 40), ["basic", "plus", "pro"][i % 3], ["الخرطوم", "جدة", "Oslo"][i % 3 if i % 5 else 0], str((i * 5) % 2)] for i in range(60)]


def table(delimiter: str) -> str:
    return "\n".join(delimiter.join(record) for record in [HEADER, *ROWS]) + "\n"


def run_cli(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, "-c", "from aidatasetkit.cli.main import main; raise SystemExit(main())", *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )


def audit(path: Path, output: Path, *extra: str) -> subprocess.CompletedProcess:
    return run_cli("audit", str(path), "--target", "churn", "--output", str(output), *extra)


def published(output: Path) -> tuple[dict, str, bytes, bytes]:
    run = read_current(output)
    manifest = (run.directory / "manifest.json").read_bytes()
    current = (output / "CURRENT").read_bytes()
    return json.loads(run.text("audit.json")), run.text("report.html"), manifest, current


def assert_nothing_published(output: Path) -> None:
    assert not (output / "CURRENT").exists()
    runs = output / "runs"
    assert not runs.exists() or not any(runs.iterdir())


@pytest.fixture(scope="module")
def semicolon(tmp_path_factory):
    folder = tmp_path_factory.mktemp("semicolon")
    path = folder / "customers.csv"
    path.write_bytes(table(";").encode("utf-8"))
    output = folder / "out"
    return audit(path, output), output, folder


class TestFilesAreReadRight:
    def test_the_semicolon_file_is_four_columns_not_one(self, semicolon):
        result, output, _ = semicolon
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        artifact, _, _, _ = published(output)
        assert artifact["ingestion"]["delimiter"] == ";"
        assert artifact["ingestion"]["delimiter_source"] == "detected"
        assert artifact["ingestion"]["column_count"] == 4
        assert artifact["ingestion"]["row_count"] == len(ROWS)
        profiled = {column["name"]["name"] for column in artifact["columns"]}
        assert profiled == set(HEADER)

    @pytest.mark.parametrize(
        ("name", "delimiter", "fmt", "source"),
        [
            ("comma.csv", ",", "csv", "detected"),
            ("tabbed.csv", "\t", "csv", "detected"),
            ("pipe.csv", "|", "csv", "detected"),
            ("plain.tsv", "\t", "tsv", "format"),
        ],
    )
    def test_every_supported_layout(self, tmp_path, name, delimiter, fmt, source):
        path = tmp_path / name
        path.write_bytes(table(delimiter).encode("utf-8"))
        result = audit(path, tmp_path / "out")
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        ingestion = published(tmp_path / "out")[0]["ingestion"]
        assert (ingestion["delimiter"], ingestion["format"], ingestion["delimiter_source"]) == (delimiter, fmt, source)

    def test_cp1256_is_read_when_named(self, tmp_path):
        path = tmp_path / "arabic.csv"
        path.write_bytes(table(",").encode("cp1256"))
        result = audit(path, tmp_path / "out", "--encoding", "cp1256")
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        artifact = published(tmp_path / "out")[0]
        assert artifact["ingestion"]["encoding"] == "cp1256"

    def test_an_explicit_tab_delimiter(self, tmp_path):
        path = tmp_path / "tabbed.csv"
        path.write_bytes(table("\t").encode("utf-8"))
        result = audit(path, tmp_path / "out", "--delimiter", "tab")
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        ingestion = published(tmp_path / "out")[0]["ingestion"]
        assert (ingestion["delimiter"], ingestion["delimiter_source"]) == ("\t", "explicit")

    @pytest.mark.parametrize(
        ("name", "text", "fmt"),
        [
            ("rows.json", json.dumps([dict(zip(HEADER, record)) for record in ROWS]), "json"),
            ("rows.jsonl", "".join(json.dumps(dict(zip(HEADER, record))) + "\n" for record in ROWS), "jsonl"),
        ],
    )
    def test_json_files_are_audited_end_to_end(self, tmp_path, name, text, fmt):
        path = tmp_path / name
        path.write_bytes(text.encode("utf-8"))
        result = audit(path, tmp_path / "out")
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        ingestion = published(tmp_path / "out")[0]["ingestion"]
        assert ingestion["format"] == fmt
        assert ingestion["row_count"] == len(ROWS)
        assert ingestion["column_count"] == 4
    def test_a_parquet_file_is_audited_end_to_end(self, tmp_path):
        pytest.importorskip("pyarrow", reason="the parquet extra is not installed")
        path = tmp_path / "rows.parquet"
        import pandas as pd

        pd.DataFrame([dict(zip(HEADER, record)) for record in ROWS]).to_parquet(path, index=False)
        result = audit(path, tmp_path / "out")
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        ingestion = published(tmp_path / "out")[0]["ingestion"]
        assert ingestion["format"] == "parquet"
        assert ingestion["row_count"] == len(ROWS)

    def test_an_xlsx_file_is_audited_end_to_end(self, tmp_path):
        pytest.importorskip("openpyxl", reason="the excel extra is not installed")
        path = tmp_path / "rows.xlsx"
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(HEADER)
        for record in ROWS:
            ws.append(record)
        wb.save(path)
        result = audit(path, tmp_path / "out")
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        ingestion = published(tmp_path / "out")[0]["ingestion"]
        assert ingestion["format"] == "xlsx"
        assert ingestion["row_count"] == len(ROWS)

    def test_a_named_sheet_is_the_one_audited(self, tmp_path):
        pytest.importorskip("openpyxl", reason="the excel extra is not installed")
        import openpyxl

        path = tmp_path / "book.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Other"
        wb.active.append(["x"])
        wb.active.append([1])
        data = wb.create_sheet("Data")
        data.append(HEADER)
        for record in ROWS[:2]:
            data.append(record)
        wb.save(path)
        result = audit(path, tmp_path / "out", "--sheet", "Data")
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        assert "2 sheets" not in result.stderr
        ingestion = published(tmp_path / "out")[0]["ingestion"]
        assert ingestion["format"] == "xlsx"
        assert ingestion["row_count"] == 2


class TestTheEvidenceSaysHowItWasRead:
    def test_the_report_has_an_input_section(self, semicolon):
        _, output, _ = semicolon
        report = published(output)[1]
        assert "<h2>Input</h2>" in report
        assert "semicolon (;)" in report and "detected" in report

    def test_no_artifact_carries_the_absolute_path(self, semicolon):
        _, output, folder = semicolon
        artifact, report, manifest, current = published(output)
        for text in (json.dumps(artifact, ensure_ascii=False), report, manifest.decode(), current.decode()):
            for form in {str(folder), str(folder.resolve()), folder.as_posix(), str(folder).replace("\\", "\\\\")}:
                assert form not in text

    def test_the_ingestion_record_holds_no_cell_value(self, semicolon):
        _, output, _ = semicolon
        record = json.dumps(published(output)[0]["ingestion"], ensure_ascii=False)
        assert "الخرطوم" not in record and "basic" not in record


class TestRefusalsPublishNothing:
    AMBIGUOUS = "x,y;z\n" + "".join(f"{i},{i};{i % 2}\n" for i in range(30))

    def test_an_ambiguous_file_is_refused_in_one_line(self, tmp_path):
        path = tmp_path / "amb.csv"
        path.write_text(self.AMBIGUOUS, encoding="utf-8")
        result = run_cli("audit", str(path), "--output", str(tmp_path / "out"))
        assert result.returncode == EXIT_CODES["usage"]
        assert result.stderr.startswith("error: AmbiguousDelimiterError:")
        assert "Traceback" not in result.stderr
        assert len(result.stderr.strip().splitlines()) == 1
        assert_nothing_published(tmp_path / "out")

    def test_the_ambiguity_is_resolved_by_naming_the_delimiter(self, tmp_path):
        path = tmp_path / "amb.csv"
        path.write_text(self.AMBIGUOUS, encoding="utf-8")
        result = run_cli("audit", str(path), "--delimiter", ";", "--output", str(tmp_path / "out"))
        assert result.returncode != EXIT_CODES["usage"], result.stderr
        assert published(tmp_path / "out")[0]["ingestion"]["column_count"] == 2

    @pytest.mark.parametrize(
        ("name", "content", "extra", "error"),
        [
            ("wrong.csv", table(",").encode("cp1256"), (), "EncodingError"),
            ("short.csv", b"a,b,churn\n1,2,0\n3,4\n", (), "MalformedInputError"),
            ("dup.csv", b"a,a,churn\n1,2,0\n", (), "DuplicateHeadersError"),
            ("header.csv", b"a,b,churn\n", (), "EmptyInputError"),
            ("semi.csv", table(";").encode(), ("--delimiter", ","), "MalformedInputError"),
        ],
    )
    def test_structured_refusals(self, tmp_path, name, content, extra, error):
        path = tmp_path / name
        path.write_bytes(content)
        result = run_cli("audit", str(path), "--output", str(tmp_path / "out"), *extra)
        assert result.returncode == EXIT_CODES["usage"]
        assert result.stderr.startswith(f"error: {error}:")
        assert "Traceback" not in result.stderr and "UnicodeDecodeError" not in result.stderr
        assert_nothing_published(tmp_path / "out")

    def test_an_unsupported_encoding_is_a_usage_error(self, tmp_path):
        path = tmp_path / "d.csv"
        path.write_bytes(table(",").encode())
        result = run_cli("audit", str(path), "--encoding", "utf-16", "--output", str(tmp_path / "out"))
        assert result.returncode != 0
        assert "--encoding" in result.stderr
        assert_nothing_published(tmp_path / "out")

    def test_a_failed_read_leaves_the_previous_run_as_it_was(self, tmp_path):
        output = tmp_path / "out"
        good = tmp_path / "good.csv"
        good.write_bytes(table(";").encode())
        assert audit(good, output).returncode != EXIT_CODES["usage"]
        before = published(output)
        runs_before = sorted(p.name for p in (output / "runs").iterdir())

        bad = tmp_path / "bad.csv"
        bad.write_bytes(b"x,y;z\n1,2;3\n4,5;6\n")
        result = run_cli("audit", str(bad), "--output", str(output))
        assert result.returncode == EXIT_CODES["usage"]
        assert published(output) == before
        assert sorted(p.name for p in (output / "runs").iterdir()) == runs_before

    def test_a_resource_refusal_leaves_the_previous_run_as_it_was(self, tmp_path):
        output = tmp_path / "out"
        good = tmp_path / "good.csv"
        good.write_bytes(table(";").encode())
        assert audit(good, output).returncode != EXIT_CODES["usage"]
        before = published(output)
        runs_before = sorted(p.name for p in (output / "runs").iterdir())

        limited = tmp_path / "limited.csv"
        limited.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
        result = audit(limited, output, "--max-rows", "1")
        assert result.returncode == EXIT_CODES["usage"]
        assert result.stderr.startswith("error: RowLimitError:")
        assert str(tmp_path) not in result.stderr and "1,2" not in result.stderr
        assert published(output) == before
        assert sorted(p.name for p in (output / "runs").iterdir()) == runs_before
